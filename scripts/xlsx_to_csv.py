#!/usr/bin/env python3
"""Opcional: gera data/pax.csv a partir da aba Base_PAX da planilha organizada.
Uso: python scripts/xlsx_to_csv.py PAX_Concorrencia_Set26_organizado.xlsx"""
import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1], data_only=True)
ws = wb["Base_PAX"]
linhas = {7: "Bondinho", 8: "AquaRio", 10: "Trem Corcovado", 11: "Paineiras", 12: "BioParque", 13: "YupStar"}
out = ["data,atracao,pax_mil"]
for col in range(2, ws.max_column + 1):
    ano, mes = ws.cell(4, col).value, ws.cell(5, col).value
    if not isinstance(ano, int):
        continue
    for r, nome in linhas.items():
        v = ws.cell(r, col).value
        if isinstance(v, (int, float)):
            out.append(f"{ano}-{mes:02d}-01,{nome},{round(v, 3)}")
open("data/pax.csv", "w", encoding="utf-8").write("\n".join(out) + "\n")
print(len(out) - 1, "linhas gravadas em data/pax.csv")
